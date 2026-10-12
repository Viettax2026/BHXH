"""Tra cứu bán tự động trang "Tra cứu hộ gia đình" của BHXH Việt Nam.

Script đọc danh sách người cần tra từ file Excel, mở trình duyệt Chrome,
tự điền từng người vào form. Người dùng chỉ cần nhìn ảnh mã xác nhận trên
trình duyệt và gõ mã đó vào cửa sổ dòng lệnh. Kết quả được ghi ra Excel
sau mỗi người, nên dừng giữa chừng vẫn giữ được phần đã tra.

Cách chạy:
    python tra_cuu_bhxh.py danh_sach.xlsx
    python tra_cuu_bhxh.py danh_sach.xlsx -o ket_qua.xlsx --bat-dau 5
    python tra_cuu_bhxh.py danh_sach.xlsx --gia-lap   (thử trên trang giả lập)
"""

import argparse
import re
import sys
import time
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright

URL = "https://baohiemxahoi.gov.vn/tracuu/Pages/tra-cuu-ho-gia-dinh.aspx"

# Tên cột trong file đầu vào (so khớp không phân biệt hoa thường, có/không dấu).
COT = {
    "tinh": ["tinh", "tinh/tp", "tinh/thanh pho", "tinh thanh"],
    "huyen": ["huyen", "quan/huyen", "quan huyen"],
    "xa": ["xa", "phuong/xa", "phuong xa"],
    "thon": ["thon", "thon/xom", "thon xom"],
    "ho_ten": ["ho ten", "ho va ten", "hoten"],
    "cccd": ["cccd", "cccd/cmnd/ho chieu", "cmnd", "so cccd", "so giay to"],
    "ngay_sinh": ["ngay sinh", "nam sinh", "ngaysinh"],
    "ma_bhxh": ["ma bhxh", "ma so bhxh", "mabhxh"],
}
TEN_COT_HIEN_THI = {
    "tinh": "Tỉnh/TP", "huyen": "Quận/Huyện", "xa": "Phường/Xã", "thon": "Thôn/Xóm",
    "ho_ten": "Họ tên", "cccd": "CCCD", "ngay_sinh": "Ngày sinh", "ma_bhxh": "Mã BHXH",
}

# Nhãn của từng ô trên trang web.
NHAN_O = {
    "tinh": "Tỉnh/TP", "huyen": "Quận/Huyện", "xa": "Phường/Xã", "thon": "Thôn/Xóm",
    "ho_ten": "Họ tên", "cccd": "CCCD/CMND", "ngay_sinh": "Ngày sinh",
    "ma_bhxh": "Mã số BHXH", "captcha": "Mã xác nhận",
}

# Tìm ô nhập (input/select) đứng ngay sau một nhãn trên trang.
JS_TIM_O = """
([nhan, loai]) => {
  const chuan = s => (s || '').normalize('NFC').replace(/\\s+/g, ' ').trim().toLowerCase();
  const can = chuan(nhan);
  const hien = el => !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
  const nhanEls = [...document.querySelectorAll('label, span, td, th, div, p, b, strong')]
    .filter(el => hien(el) && chuan(el.textContent).startsWith(can) && chuan(el.textContent).length < can.length + 15);
  if (!nhanEls.length) return null;
  // Lấy phần tử nhãn nhỏ nhất (không chứa phần tử nhãn khác).
  const nhanEl = nhanEls.find(el => !nhanEls.some(o => o !== el && el.contains(o))) || nhanEls[0];
  if (nhanEl.tagName === 'LABEL' && nhanEl.htmlFor) {
    const o = document.getElementById(nhanEl.htmlFor);
    if (o && (!loai || o.tagName.toLowerCase() === loai)) return o.id ? '#' + CSS.escape(o.id) : null;
  }
  const sel = loai ? loai : 'input:not([type=hidden]):not([type=radio]):not([type=checkbox]):not([type=submit]):not([type=button]):not([type=image]), select, textarea';
  const ung = [...document.querySelectorAll(sel)].filter(hien);
  const o = ung.find(el => nhanEl.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  if (!o) return null;
  if (!o.id) o.id = 'tcbhxh_' + Math.random().toString(36).slice(2);
  return '#' + CSS.escape(o.id);
}
"""

JS_LAY_BANG = """
() => [...document.querySelectorAll('table')]
  .filter(t => t.offsetWidth || t.offsetHeight)
  .filter(t => !t.querySelector('table, input, select'))
  .map(t => [...t.rows].map(r => [...r.cells].map(c => c.innerText.replace(/\\s+/g, ' ').trim())))
  .filter(rows => rows.length > 0)
"""


def bo_dau(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("đ", "d").replace("Đ", "D").lower().strip()


def chuan_hoa(s):
    return re.sub(r"\s+", " ", bo_dau(s)).strip()


def gia_tri_o(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def doc_dau_vao(duong_dan):
    wb = load_workbook(duong_dan, data_only=True)
    ws = wb.active
    hang = list(ws.iter_rows(values_only=True))
    if not hang:
        sys.exit("File đầu vào trống.")
    tieu_de = [chuan_hoa(h) for h in hang[0]]
    vi_tri = {}
    for khoa, ten_cho_phep in COT.items():
        for i, h in enumerate(tieu_de):
            if h in ten_cho_phep:
                vi_tri[khoa] = i
                break
    thieu = [TEN_COT_HIEN_THI[k] for k in ("tinh", "ho_ten", "cccd") if k not in vi_tri]
    if thieu:
        sys.exit("File đầu vào thiếu cột bắt buộc: " + ", ".join(thieu))
    nguoi = []
    for so_hang, h in enumerate(hang[1:], start=2):
        ban_ghi = {k: gia_tri_o(h[i]) if i < len(h) else "" for k, i in vi_tri.items()}
        if not any(ban_ghi.values()):
            continue
        ban_ghi["_hang"] = so_hang
        nguoi.append(ban_ghi)
    return nguoi


class GhiKetQua:
    def __init__(self, duong_dan):
        self.duong_dan = duong_dan
        self.wb = Workbook()
        self.tong_hop = self.wb.active
        self.tong_hop.title = "KetQua"
        self.tong_hop.append(["Hàng nguồn"] + [TEN_COT_HIEN_THI[k] for k in COT]
                             + ["Trạng thái", "Thông báo", "Kết quả (văn bản)", "Thời gian tra"])
        self.chi_tiet = self.wb.create_sheet("ChiTiet")
        self.chi_tiet.append(["Hàng nguồn", "Họ tên tra", "CCCD tra", "Bảng số", "Dữ liệu hàng →"])
        for ws in (self.tong_hop, self.chi_tiet):
            for c in ws[1]:
                c.font = Font(bold=True)

    def them(self, nguoi, trang_thai, thong_bao, van_ban, bang):
        self.tong_hop.append([nguoi["_hang"]] + [nguoi.get(k, "") for k in COT]
                             + [trang_thai, thong_bao, van_ban[:32000],
                                datetime.now().strftime("%d/%m/%Y %H:%M:%S")])
        for so_bang, rows in enumerate(bang, start=1):
            for r in rows:
                self.chi_tiet.append([nguoi["_hang"], nguoi.get("ho_ten", ""), nguoi.get("cccd", ""), so_bang] + r)
        self.wb.save(self.duong_dan)


class TrangTraCuu:
    def __init__(self, page, url):
        self.page = page
        self.url = url
        self.thong_bao_dialog = []
        page.on("dialog", self._khi_co_dialog)

    def _khi_co_dialog(self, dialog):
        self.thong_bao_dialog.append(dialog.message)
        dialog.accept()

    def tim_o(self, khoa, loai=None):
        return self.page.evaluate(JS_TIM_O, [NHAN_O[khoa], loai])

    def chon_dropdown(self, khoa, gia_tri):
        if not gia_tri:
            return True
        sel = self.tim_o(khoa, "select")
        if not sel:
            print(f"  ! Không tìm thấy ô {NHAN_O[khoa]}, hãy chọn tay trên trình duyệt.")
            return False
        can = chuan_hoa(gia_tri)
        # Đợi danh sách được nạp (huyện/xã/thôn nạp sau khi chọn cấp trên).
        lua_chon = []
        for _ in range(40):
            lua_chon = self.page.eval_on_selector(
                sel, "s => [...s.options].map(o => ({v: o.value, t: o.text}))")
            if len(lua_chon) > 1:
                break
            time.sleep(0.25)
        khop = [o for o in lua_chon if chuan_hoa(o["t"]) == can]
        if not khop:
            # Cho phép bỏ tiền tố: "Hà Nội" khớp "Thành phố Hà Nội".
            khop = [o for o in lua_chon if chuan_hoa(o["t"]).endswith(" " + can)]
        if not khop:
            print(f"  ! Không thấy '{gia_tri}' trong {NHAN_O[khoa]}, hãy chọn tay trên trình duyệt.")
            return False
        self.page.select_option(sel, value=khop[0]["v"])
        self.page.wait_for_load_state("networkidle")
        time.sleep(0.5)
        return True

    def dien(self, khoa, gia_tri):
        if not gia_tri:
            return
        sel = self.tim_o(khoa)
        if not sel:
            print(f"  ! Không tìm thấy ô {NHAN_O[khoa]}, hãy nhập tay trên trình duyệt.")
            return
        self.page.fill(sel, gia_tri)

    def chon_kieu_ten(self, ho_ten):
        co_dau = bo_dau(ho_ten) != ho_ten.lower().strip()
        nhan = "Có dấu" if co_dau else "Không dấu"
        o = self.page.get_by_text(nhan, exact=True)
        if o.count():
            o.first.click()

    def lam_moi_captcha(self):
        anh = self.page.locator("img[src*='Captcha' i], img[id*='Captcha' i], img[title*='đổi mã']")
        if anh.count():
            anh.first.click()
            time.sleep(0.5)

    def dien_form(self, nguoi):
        self.page.goto(self.url, wait_until="networkidle")
        for khoa in ("tinh", "huyen", "xa", "thon"):
            self.chon_dropdown(khoa, nguoi.get(khoa))
        self.dien("ho_ten", nguoi.get("ho_ten"))
        self.chon_kieu_ten(nguoi.get("ho_ten", ""))
        self.dien("cccd", nguoi.get("cccd"))
        self.dien("ngay_sinh", nguoi.get("ngay_sinh"))
        self.dien("ma_bhxh", nguoi.get("ma_bhxh"))

    def tra_cuu(self, ma_xac_nhan):
        sel = self.tim_o("captcha")
        if sel:
            self.page.fill(sel, ma_xac_nhan)
        bang_truoc = self.page.evaluate(JS_LAY_BANG)
        dong_truoc = set(self.page.inner_text("body").splitlines())
        self.thong_bao_dialog.clear()
        nut = self.page.get_by_text("Tra cứu", exact=True)
        if not nut.count():
            nut = self.page.locator("input[value='Tra cứu'], button:has-text('Tra cứu')")
        nut.last.click()
        try:
            self.page.wait_for_load_state("networkidle", timeout=20000)
        except Exception:
            pass
        time.sleep(1.5)
        bang = [b for b in self.page.evaluate(JS_LAY_BANG) if b not in bang_truoc]
        van_ban = "\n".join(" | ".join(c for c in r if c) for b in bang for r in b)
        # Thông báo mới hiện trên trang (vd: sai mã xác nhận, không tìm thấy).
        dong_moi = [d.strip() for d in self.page.inner_text("body").splitlines()
                    if d.strip() and "\t" not in d and d not in dong_truoc and d.strip() not in van_ban]
        thong_bao = self.thong_bao_dialog + dong_moi[:5]
        return bang, van_ban, "; ".join(thong_bao)


def main():
    ap = argparse.ArgumentParser(description="Tra cứu bán tự động hộ gia đình BHXH từ file Excel.")
    ap.add_argument("dau_vao", help="File Excel danh sách người cần tra")
    ap.add_argument("-o", "--dau-ra", help="File Excel kết quả (mặc định: ket_qua_<thời gian>.xlsx)")
    ap.add_argument("--bat-dau", type=int, default=1, help="Bắt đầu từ người thứ mấy trong danh sách (mặc định 1)")
    ap.add_argument("--gia-lap", action="store_true", help="Thử trên trang giả lập gia_lap.html thay vì trang BHXH thật")
    ap.add_argument("--url", help="Địa chỉ trang tra cứu khác (mặc định: trang BHXH thật)")
    args = ap.parse_args()
    url = args.url or (Path(__file__).with_name("gia_lap.html").resolve().as_uri() if args.gia_lap else URL)

    nguoi = doc_dau_vao(args.dau_vao)[args.bat_dau - 1:]
    dau_ra = args.dau_ra or f"ket_qua_{datetime.now():%Y%m%d_%H%M}.xlsx"
    ghi = GhiKetQua(dau_ra)
    if url != URL:
        print(f"Đang dùng trang: {url}")
    print(f"Có {len(nguoi)} người cần tra. Kết quả ghi vào: {Path(dau_ra).resolve()}")
    print("Với mỗi người: nhìn ảnh mã xác nhận trên trình duyệt rồi gõ mã vào đây.")
    print("  Gõ 'r' để đổi ảnh mã khác, 's' để bỏ qua người này, 'q' để dừng.\n")

    with sync_playwright() as p:
        try:
            trinh_duyet = p.chromium.launch(channel="chrome", headless=False)
        except Exception:
            trinh_duyet = p.chromium.launch(headless=False)
        page = trinh_duyet.new_page()
        trang = TrangTraCuu(page, url)

        for i, ng in enumerate(nguoi, start=1):
            print(f"[{i}/{len(nguoi)}] {ng.get('ho_ten')} - {ng.get('cccd')}")
            trang.dien_form(ng)
            while True:
                ma = input("  Mã xác nhận: ").strip()
                if ma.lower() == "q":
                    trinh_duyet.close()
                    print(f"Đã dừng. Kết quả lưu tại {dau_ra}")
                    return
                if ma.lower() == "s":
                    ghi.them(ng, "Bỏ qua", "", "", [])
                    break
                if ma.lower() == "r" or not ma:
                    trang.lam_moi_captcha()
                    continue
                bang, van_ban, thong_bao = trang.tra_cuu(ma)
                if thong_bao:
                    print(f"  Thông báo: {thong_bao}")
                print("  Kết quả: " + (van_ban[:300] or "(không thấy bảng kết quả)"))
                tiep = input("  Enter = lưu và sang người tiếp, 'l' = tra lại người này: ").strip().lower()
                if tiep == "l":
                    trang.dien_form(ng)
                    continue
                trang_thai = "Có kết quả" if bang else "Không có kết quả"
                ghi.them(ng, trang_thai, thong_bao, van_ban, bang)
                break
        trinh_duyet.close()
    print(f"\nXong. Kết quả lưu tại {Path(dau_ra).resolve()}")


if __name__ == "__main__":
    main()
